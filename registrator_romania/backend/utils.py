from datetime import datetime
from functools import partial
import asyncio
import multiprocessing
import threading
from pathlib import Path
import random
import re
import string
import sys
from zoneinfo import ZoneInfo
import aiofiles
import dateutil
from docx import Document
from loguru import logger
import openpyxl
import pandas as pd

from registrator_romania.shared import get_config


def divide_list(src_list: list, divides: int = 100):
    return [src_list[x : x + divides] for x in range(0, len(src_list), divides)]


def filter_by_log_level(loglevels: list[str]):
    return lambda record: record["level"].name in loglevels


def is_host_port(v: str):
    if re.findall(r"\d+:\d+", v):
        return True


class Transliterator:
    def __init__(self, language):
        if language == "tr":
            self.translit_dict = {
                "Ş": "S",
                "ş": "s",
                "İ": "I",
                "ı": "i",
                "Ğ": "G",
                "ğ": "g",
                "Ç": "C",
                "ç": "c",
                "Ö": "O",
                "ö": "o",
                "Ü": "U",
                "ü": "u",
            }
        else:
            self.translit_dict = {}

    def transliterate(self, text):
        return "".join(self.translit_dict.get(char, char) for char in text)


def prepare_users_data(users_data: list[dict]):
    keys = [
        "Prenume Pasaport",
        "Nume Pasaport",
        "Data nasterii",
        "Locul naşterii",
        "Prenume Mama",
        "Prenume Tata",
        "Adresa de email",
        "Serie și număr Pașaport",
    ]
    objs = []
    for us_data in users_data:
        obj = {}
        for k, v in us_data.items():
            if len(str(v)) < 3:
                # Append to value last symbol while len(v) < 3
                while len(str(v)) < 3:
                    v += v[-1]

            # Replace values like `Doğum tarihi:09.09.1976`
            v = v.split(":")[-1].strip()
            # Change turkey letters on english letters
            v = Transliterator("tr").transliterate(v)

            if k == "Data nasterii":
                try:
                    dt = dateutil.parser.parse(v, dayfirst=False)
                except dateutil.parser.ParserError:
                    dt = dateutil.parser.parse(v, dayfirst=True)

                # We need to format date like `1976-09-09`
                v = dt.strftime("%Y-%m-%d")
                assert datetime.strptime(v, "%Y-%m-%d")

            obj[k] = v

        # Tranform case
        obj["Nume Pasaport"] = obj["Nume Pasaport"].upper()
        obj["Prenume Pasaport"] = obj["Prenume Pasaport"].upper()
        obj["Adresa de email"] = obj["Adresa de email"].lower()
        objs.append(obj)

    assert all(k in obj for k in keys for obj in objs)
    return objs


def get_users_data_from_docx():
    doc = Document("users.docx")

    users = []
    user = {}
    mapping = {
        "Prenume Pasaport": ["Prenume"],
        "Nume Pasaport": ["Nume"],
        "Data nasterii": ["Data naşterii"],
        "Locul naşterii": ["Locul naşterii"],
        "Prenume Mama": ["Prenumele mamei", "Numele mame", "Numele mamei"],
        "Prenume Tata": [
            "Prenumele tatalui",
            "Numele tatalui",
        ],
        "Adresa de email": ["Adresa de e-mail"],
        "Serie și număr Pașaport": ["Seria şi numar Paşaport"],
    }
    for paragraph in doc.paragraphs:
        text = paragraph.text.replace("\n", "").strip()
        if not text:
            continue

        record = re.findall(r"(^[\d\.]*)(.*)", text)[0][1]
        col, val = list(map(lambda v: v.strip(), record.split(":")))

        key = None
        for k, v in mapping.items():
            if col in v:
                key = k
                break

        assert key
        val = Transliterator("tr").transliterate(val)

        if key == "Data nasterii":
            try:
                dt = datetime.strptime(val, "%Y-%m-%d")
            except Exception:
                dt = datetime.strptime(val, "%d-%m-%Y")
            val = dt.strftime("%Y-%m-%d")
        elif key == "":
            val = val.lower()

        user[key] = val

        if key == "Serie și număr Pașaport":
            users.append(user.copy())
            user.clear()

    return prepare_users_data(users)


def get_users_data_from_csv(file_path: str = None):
    if file_path:
        df = pd.read_csv(file_path)
    else:
        df = pd.read_csv("users.csv")

    users_data = df.to_dict("records")
    return prepare_users_data(users_data)


async def async_log_message(message: str, path: Path = None):
    if path:
        async with aiofiles.open(path, mode="a", encoding="utf-8") as f:
            await f.write(message)

def setup_loggers(registration_date: datetime, save_logs: bool = True):
    dirpath = f"registrations_{registration_date.strftime("%d.%m.%Y")}"
    Path(dirpath).mkdir(parents=True, exist_ok=True)

    if save_logs:
        logger.remove()
        logger.add(
            sys.stderr,
            filter=filter_by_log_level(loglevels=["INFO", "SUCCESS", "ERROR"]),
            # enqueue=True,
            backtrace=True,
            diagnose=True,
        )
        logger.add(
            Path().joinpath(dirpath, "errors.log"),
            # partial(async_log_message, path=Path().joinpath(dirpath, "errors.log")),
            filter=filter_by_log_level(loglevels=["ERROR"]),
            enqueue=True,
            backtrace=True,
            diagnose=True,
        )
        logger.add(
            Path().joinpath(dirpath, "debug.log"),
            # partial(async_log_message, path=Path().joinpath(dirpath, "debug.log")),
            filter=filter_by_log_level(loglevels=["DEBUG"]),
            enqueue=True,
            backtrace=True,
            diagnose=True,
        )
        logger.add(
            Path().joinpath(dirpath, "success.log"),
            # partial(async_log_message, path=Path().joinpath(dirpath, "success.log")),
            filter=filter_by_log_level(loglevels=["SUCCESS"]),
            enqueue=True,
            backtrace=True,
            diagnose=True,
        )


def get_users_data_from_txt():
    keys = [
        "Prenume Pasaport",
        "Nume Pasaport",
        "Data nasterii",
        "Locul naşterii",
        "Prenume Mama",
        "Prenume Tata",
        "Adresa de email",
        "Serie și număr Pașaport",
    ]
    with open("users.txt") as f:
        data = f.read()

    objs = []
    for values in data.split("\n\n"):
        obj = {}
        for k, v in zip(keys, values.split("\n")):
            v = v.split(":")[-1].strip()
            v = Transliterator("tr").transliterate(v)
            if k == "Data nasterii":
                try:
                    dt = datetime.strptime(v, "%Y-%m-%d")
                except Exception:
                    dt = datetime.strptime(v, "%d-%m-%Y")
                v = dt.strftime("%Y-%m-%d")

            obj[k] = v
        obj["Adresa de email"] = obj["Adresa de email"].lower()
        objs.append(obj)

    return prepare_users_data(objs)


def get_users_data_from_xslx(path: str = None):
    keys = [
        "Prenume Pasaport",
        "Nume Pasaport",
        "Data nasterii",
        "Locul naşterii",
        "Prenume Mama",
        "Prenume Tata",
        "Adresa de email",
        "Serie și număr Pașaport",
    ]

    w = openpyxl.load_workbook("users.xlsx" if not path else path)
    sheet = w.active
    data = []
    for row in sheet.iter_rows(min_row=0, max_row=None, values_only=True):
        if not data:
            assert row
            row = keys.copy()
        row = list(row)

        if row and any(not isinstance(v, str) for v in row):
            for v in row:
                if isinstance(v, datetime):
                    row[row.index(v)] = v.strftime("%Y-%m-%d")

        data.append(row)
    df = pd.DataFrame(data)
    df.columns = df.iloc[0]
    df.drop(df.index[0], inplace=True)
    objs_raw = df.to_dict("records")
    return prepare_users_data(objs_raw)


def generate_fake_users_data(n: int = 20):
    def random_string(n: int = 5):
        return "".join(random.choice(string.ascii_uppercase) for _ in range(n))

    return [
        {
            "Nume Pasaport": f"ARNA{random_string(2)}",
            "Prenume Pasaport": f"VALE{random_string(2)}",
            "Data nasterii": f"199{random.randint(0, 9)}-10-1{random.randint(0, 9)}",
            "Locul naşterii": "ISTANBUL",
            "Prenume Mama": "RECYE",
            "Prenume Tata": "SABRI",
            "Adresa de email": f"{random_string(7)}@gmail.com",
            "Serie și număr Pașaport": f"U{random.randint(10_000_000, 10_999_999)}",
        }
        for _ in range(n)
    ]

def get_dt_moscow() -> datetime:
    return datetime.now().astimezone(ZoneInfo("Europe/Moscow"))


def get_success_regs_time_from_debug_log(content_log: str) -> list[datetime]:
    lines = content_log.splitlines()

    def find_first_request(task_id):
        for line in lines:
            if line.count("send request") and line.count(task_id):
                time = re.search(r"^([^|]+) |", line).group(1)
                dt = datetime.strptime(time, "%Y-%m-%d %H:%M:%S.%f")
                return dt

    data = []
    uniq = []
    for line in lines:
        if line.count("success: True"):
            task_id = re.search(r"Task.*('Task-\d+')", line).group(1)
            time = find_first_request(task_id=task_id)
            data.append(time)
            if time not in uniq:
                uniq.append(time)

    return {"all": sorted(data), "uniq": sorted(uniq)}


def get_requests_times_from_log(log_content: str):
    lines = log_content.splitlines()
    data = []

    for line in lines:
        if line.count("send request"):
            endpoint = (
                re.search(r"send request on (https{0,1}://[^ ]+).*proxy", line)
                .group(1)
                .rstrip(".")
            )
            strtime = re.search(r"^([^|]+) |", line).group(1)

            obj = {
                "time": datetime.strptime(strtime, "%Y-%m-%d %H:%M:%S.%f"),
                "endpoint": endpoint,
            }
            data.append(obj)

    return sorted(data, key=lambda x: x["time"])


def get_rpc_times(log_content: str):
    times = get_requests_times_from_log(log_content)
    second = None
    prev_time = None

    results = []

    obj = {}
    for t in times:
        endpoint = t["endpoint"]
        if not obj.get(endpoint):
            obj[endpoint] = 0

        if second is None or (t["time"] - prev_time).seconds >= 1:
            if second:
                obj["time"] = str(prev_time)
                results.append(obj)
                obj = {endpoint: 0}

            prev_time = t["time"]
            second = prev_time.second

        obj[endpoint] += 1

    return sorted(results, key=lambda x: x["time"])


def get_current_info():
    atask = asyncio.current_task()
    return (
        f"{datetime.now()} - "
        f"[Process: {multiprocessing.current_process().name}] "
        f"[Thread: {threading.current_thread().name}] "
        f"[Task: '{atask.get_name()}': {atask.get_coro().__qualname__}]"
    )


from pprint import pprint


async def main():
    from concurrent.futures import ThreadPoolExecutor
    from bindings2 import test_request, CaptchaPasser, APIRomania
    # print(test_request())
    users_data = generate_fake_users_data(1)
    
    # passer = CaptchaPasser()
    reg_date = "2025-01-14"
    api = APIRomania()

    def registrate(user_data):
        # token = passer.get_recaptcha_token()
        res = api.make_registration(user_data, 2, reg_date, proxy="http://xP3SrZGCP3C6:RNW78Fm5@pool.proxy.market:10005")
        success = isinstance(res, str) and res.count("Felicitări")
        print(f"user num: {users_data.index(user_data)}, success: {bool(success)}")
        
    
    with ThreadPoolExecutor() as pool:
        start = datetime.now()
        for res in pool.map(registrate, users_data):
            ...
        print(datetime.now() - start)
        
    ...
    
    
if __name__ == "__main__":
    asyncio.run(main())
    # c = open("registrations_02.12.2024/debug.log").read()
    # # c = open("logs.log").read()
    # # c = open("/home/daniil/Downloads/Telegram Desktop/debug (22).log").read()

    # times = get_success_regs_time_from_debug_log(c)
    # all_req = get_requests_times_from_log(c)
    # # times = get_rpc_times(c)

    # # pprint(times)
    # pprint(times["uniq"])
    # pprint(len(all_req))
    # print(f"{len(times['uniq'])}/{len(times['all'])}")
    # a = 'uniq'
    # print((times[a][-1] - times[a][0]).total_seconds())
    # print(len(times))
